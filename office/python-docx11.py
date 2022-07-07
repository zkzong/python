from docx import Document
import psutil

#准备一些数据
vmem = psutil.virtual_memory()
vmem_dict = vmem._asdict()

trow = 2
tcol = len(vmem_dict.keys())
#产生表格

document = Document()
table = document.add_table(rows=trow,cols=tcol,style = 'Colorful Grid Accent 4')
for col,info in enumerate(vmem_dict.keys()):
    table.cell(0,col).text = info
    if info == 'percent':
        table.cell(1,col).text = str(vmem_dict[info])+'%'
    else:
        table.cell(1,col).text = str(vmem_dict[info]/(1024*1024)) + 'M'
document.save('table.docx')